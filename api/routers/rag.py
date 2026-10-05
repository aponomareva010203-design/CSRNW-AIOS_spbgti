"""HTTP-слой RAG-баз (api/routers/rag.py).

GET    /api/rag                        — RAG-базы текущего пользователя + базы,
                                          привязанные к проектам, где
                                          пользователь участник (Просмотр+);
GET    /api/rag/{id}                   — одна RAG-база (своя или с доступом);
PATCH  /api/rag/{id}                   — правка карточки (частичный словарь
                                          полей): владелец или участник
                                          проекта с «Редактированием»;
POST   /api/rag/{id}/publish           — публикация в каталог «Данные и знания»:
                                          владелец или участник проекта
                                          с «Полным доступом»;
POST   /api/rag/{id}/documents         — загрузка документа (multipart), ставит
                                          статус queued и запускает индексацию
                                          фоном (правка и выше);
GET    /api/rag/{id}/documents         — статусы документов базы (для страницы
                                          загрузки — опрос без повторных полных
                                          карточек);
GET    /api/rag/{id}/ontology-check    — сверка с текущей опубликованной версией
                                          привязанной онтологии: что изменилось в
                                          составе классов и сколько чанков затронет
                                          актуализация (см. §2.15 архитектуры);
POST   /api/rag/{id}/reclassify        — актуализация class_ids чанков под текущую
                                          онтологию (без пере-эмбеддинга, с бэкапом
                                          старых значений — см. scripts/backfill_class_ids.py
                                          для того же механизма без HTTP);
DELETE /api/rag/{id}                   — удаление, только роль «Администратор»
                                          (любой RAG-базы по id, вместе с публикацией).

RAG-базы конструируются навыком «Конструктор RAG» через инструменты get_rag /
save_rag в рабочем чате (конфигурация — тема, ограничения, типовые вопросы);
сами документы загружаются здесь, на отдельной странице (data/skills/
rag-builder.md направляет туда пользователя), не через чат. Индексация —
текст → чанки → эмбеддинги → Qdrant (runtime/indexer.py) — идёт в фоне
(BackgroundTasks), пользователь опрашивает статус через GET .../documents.
Собственные RAG-базы адресуются по идентификатору владельца из Bearer-токена
(get_current_user) — из запроса не принимается, чужие недостижимы. База,
привязанная к проекту (project_id), открывается участникам проекта с уровнем
Просмотр и выше (runtime/projects.LEVELS — названия уровней как на GitHub),
правка — с «Редактированием», публикация — с «Полным доступом».
"""

import io
import json
import re
import secrets
import time
import zipfile
from pathlib import Path

from qdrant_client.http import models as qmodels

from fastapi import APIRouter, BackgroundTasks, Body, Depends, File, HTTPException, UploadFile

from ..runtime import indexer, ontologies, projects, rag, team
from ..runtime.rag import RagError
from ..runtime.users import ROLE_ADMIN
from .auth import get_current_user, require_admin

router = APIRouter(prefix="/api")

DOC_ID_RE = re.compile(r"^doc-[a-f0-9]{8,12}$")
ZIP_MAX_BYTES = 100 * 1024 * 1024   # сам архив
ZIP_MAX_TOTAL = 300 * 1024 * 1024   # суммарно после распаковки (защита от zip-бомбы)
ZIP_MAX_FILES = 200                 # файлов в архиве
DOC_EXTS = (".pdf", ".txt", ".md", ".docx")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 МБ на файл — разумный предел для регламентов/статей


def _resolve(store, rag_id: str, user: dict, need: str, team_kind: str | None = None):
    """Карточка RAG-базы и идентификатор владельца для действий по уровню.

    Своя база — (card, user["id"]). Привязанная к проекту база отдаётся
    участнику с уровнем need и выше (см. projects.level_allows); для правки
    засчитывается и старое право «RAG-системы» из «Моей команды»
    (team_kind="rag"). Админ роли — всё. Остальным — (None, None): чужая
    база неотличима от несуществующей.
    """
    card = store.get(user["id"], rag_id)
    if card is not None:
        return card, user["id"]
    found = store.find(rag_id)
    if found is None or (found.get("admin_only") and user.get("role") != ROLE_ADMIN):
        return None, None
    if user.get("role") == ROLE_ADMIN:
        return found, found["user_id"]
    level = projects.artifact_level(found.get("project_id") or "", user)
    if projects.level_allows(level, need):
        return found, found["user_id"]
    if team_kind and team.get_store().can_edit(found["user_id"], user["id"], team_kind):
        return found, found["user_id"]
    return None, None


def _new_doc_id() -> str:
    return "doc-" + secrets.token_hex(5)


def _extract_text(filename: str, content: bytes) -> str:
    """Извлекает текст из загруженного файла; поддерживает .txt/.md и .pdf."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        import io

        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if lower.endswith(".docx"):
        try:
            import docx
        except ImportError:
            raise RuntimeError("на сервере не установлен python-docx")
        import io as _io

        d = docx.Document(_io.BytesIO(content))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return content.decode("utf-8", errors="ignore")


@router.get("/rag")
def list_rag(user: dict = Depends(get_current_user)):
    """RAG-базы владельца токена + привязанные к его проектам, свежие сверху."""
    store = rag.get_store()
    is_admin = user.get("role") == ROLE_ADMIN
    items: list = []
    seen: set = set()
    for card in store.list(user["id"]):
        seen.add(card["id"])
        items.append(store.summary(card))
    for card in store.list_all():
        if card["id"] in seen or not card.get("project_id"):
            continue
        if card.get("admin_only") and not is_admin:
            continue
        level = projects.artifact_level(card.get("project_id") or "", user)
        if not projects.level_allows(level, projects.LEVEL_READ):
            continue
        view = store.summary(card)
        view["shared"] = True
        view["my_level"] = level
        items.append(view)
    items.sort(key=lambda c: c.get("updated_at") or 0, reverse=True)
    return {"items": items}


@router.get("/rag/{rag_id}")
def get_rag_card(rag_id: str, user: dict = Depends(get_current_user)):
    """Одна RAG-база (полная карточка): своя или привязанная к проекту
    с доступом (флаг shared, уровень my_level)."""
    store = rag.get_store()
    card, _owner_id = _resolve(store, rag_id, user, projects.LEVEL_READ)
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    view = store.public_card(card)
    if card.get("user_id") != user["id"]:
        view["shared"] = True
        view["my_level"] = projects.artifact_level(card.get("project_id") or "", user)
    return view


@router.patch("/rag/{rag_id}")
def update_rag(
    rag_id: str,
    patch: dict = Body(default={}),
    user: dict = Depends(get_current_user),
):
    """Правка карточки владельцем токена.

    Тело — частичный словарь полей (name, domain, description, author, tags,
    project_id, ontology_id, corpus_types, corpus_volume, corpus_format,
    access_restrictions, sensitive, typical_queries, citation_required,
    owner, update_frequency, status): переданные поля перезаписываются,
    остальные сохраняются.

    Чужая база доступна участнику проекта с правом «Редактирование» либо
    участнику команды владельца с правом правки RAG-систем («Моя команда»;
    удалять её может только администратор). Тестовая база (admin_only)
    правится только администраторами.
    """
    store = rag.get_store()
    card, owner_id = _resolve(store, rag_id, user, projects.LEVEL_WRITE, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    try:
        saved = store.save(owner_id, dict(patch or {}, rag_id=rag_id))
    except RagError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    return store.public_card(saved)


@router.post("/rag/{rag_id}/publish")
def publish_rag(rag_id: str, user: dict = Depends(get_current_user)):
    """Публикация RAG-базы в каталог «Данные и знания».

    Публикует владелец токена или участник проекта с «Полным доступом»
    (Просмотр/Редактирование публикацию не дают — она открывает базу всем
    посетителям сайта). Проверяет полноту (валидация require_complete:
    нужен минимум один документ в статусе ready), переводит в active и
    выкладывает карточку в общий каталог; идемпотентно. Открытый доступ
    на чтение (модель А): после публикации базу может подключить в чате
    любой пользователь.
    """
    store = rag.get_store()
    card, owner_id = _resolve(store, rag_id, user, projects.LEVEL_ADMIN)
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    try:
        saved = store.save(owner_id, {"rag_id": rag_id, "status": "active"})
    except RagError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    return {
        "ok": True,
        "rag": store.public_card(saved),
        "catalog_entry": store.catalog_entry(saved),
    }


def _zip_name(info) -> str:
    # имена без флага UTF-8 zipfile читает как cp437: возвращаем байты и пробуем utf-8, затем cp866 (Windows)
    name = info.filename
    if info.flag_bits & 0x800:
        return name
    for enc in ("utf-8", "cp866"):
        try:
            return name.encode("cp437").decode(enc)
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
    return name


def _plan_zip(content: bytes):
    # читает только оглавление архива; возвращает (план, пропущенные)
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="Не удалось открыть архив: повреждённый .zip")
    with zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
    if len(infos) > ZIP_MAX_FILES:
        raise HTTPException(status_code=413, detail=f"В архиве больше {ZIP_MAX_FILES} файлов")
    if sum(i.file_size for i in infos) > ZIP_MAX_TOTAL:
        raise HTTPException(status_code=413, detail="Архив слишком большой после распаковки (лимит 300 МБ)")
    plan, skipped = [], []
    for info in infos:
        full = _zip_name(info).replace("\\", "/")
        parts = [x for x in full.split("/") if x not in ("", ".", "..")]
        if not parts or any(x.startswith(".") or x == "__MACOSX" for x in parts):
            continue  # служебные файлы молча пропускаем
        shown = "/".join(parts)
        if info.flag_bits & 0x1:
            skipped.append({"name": shown, "reason": "архив зашифрован паролем"})
        elif not parts[-1].lower().endswith(DOC_EXTS):
            skipped.append({"name": shown, "reason": "неверный формат"})
        elif info.file_size > MAX_UPLOAD_BYTES:
            skipped.append({"name": shown, "reason": "превышен лимит 10 МБ"})
        else:
            plan.append({"raw": info.filename, "folder": "/".join(parts[:-1]), "filename": parts[-1]})
    return plan, skipped


@router.post("/rag/{rag_id}/documents", status_code=201)
async def upload_document(
    rag_id: str,
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    """Загружает документ в базу, ставит в очередь (правка и выше).

    Отвечает сразу (статус queued), извлечение текста и индексация идут
    в фоне (BackgroundTasks) — не блокируют HTTP-ответ. Прогресс опрашивать
    через GET /api/rag/{id}/documents.
    """
    store = rag.get_store()
    card, owner_id = _resolve(store, rag_id, user, projects.LEVEL_WRITE, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")

    content = await file.read()
    filename = file.filename or "документ"
    is_zip = filename.lower().endswith(".zip")
    if len(content) > (ZIP_MAX_BYTES if is_zip else MAX_UPLOAD_BYTES):
        raise HTTPException(status_code=413, detail="Архив больше 100 МБ" if is_zip else "Файл больше 10 МБ")

    if is_zip:
        plan, skipped = _plan_zip(content)
        if not plan:
            raise HTTPException(status_code=422, detail="В архиве нет файлов .pdf, .txt, .md, .docx")
        new_docs = [
            {"doc_id": _new_doc_id(), "filename": pl["filename"][-150:], "folder": pl["folder"], "status": "queued"}
            for pl in plan
        ]
        try:
            store.save(owner_id, {**card, "documents": list(card.get("documents") or []) + new_docs})
        except RagError as exc:
            raise HTTPException(status_code=422, detail=str(exc))

        async def _process_zip() -> None:
            total = 0
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                for pl, d in zip(plan, new_docs):
                    try:
                        with zf.open(zf.getinfo(pl["raw"])) as fh:
                            data = fh.read(MAX_UPLOAD_BYTES + 1)
                        total += len(data)
                        if len(data) > MAX_UPLOAD_BYTES or total > ZIP_MAX_TOTAL:
                            raise ValueError("превышен лимит размера при распаковке")
                        text = _extract_text(d["filename"], data)
                    except Exception as exc:
                        store.update_document(
                            owner_id, rag_id, d["doc_id"],
                            {"status": rag.DOC_STATUS_ERROR, "error": f"Не удалось извлечь текст: {exc}"[:500]},
                        )
                        continue
                    await indexer.index_document(
                        owner_id, rag_id, d["doc_id"], d["filename"], text,
                        ontology_id=card.get("ontology_id") or None,
                    )

        background_tasks.add_task(_process_zip)
        return {"ok": True, "archive": True, "queued": len(new_docs), "skipped": skipped}

    doc_id = _new_doc_id()
    documents = list(card.get("documents") or [])
    documents.append({"doc_id": doc_id, "filename": filename, "status": "queued"})
    try:
        store.save(owner_id, {**card, "documents": documents})
    except RagError as exc:
        raise HTTPException(status_code=422, detail=str(exc))

    async def _process() -> None:
        try:
            text = _extract_text(filename, content)
        except Exception as exc:  # извлечение текста — внешний код (pypdf), ловим широко
            store.update_document(
                owner_id, rag_id, doc_id,
                {"status": rag.DOC_STATUS_ERROR, "error": f"Не удалось извлечь текст: {exc}"[:500]},
            )
            return
        await indexer.index_document(
            owner_id, rag_id, doc_id, filename, text,
            ontology_id=card.get("ontology_id") or None,
        )

    background_tasks.add_task(_process)
    return {"ok": True, "doc_id": doc_id, "status": "queued"}


@router.get("/rag/{rag_id}/documents")
def list_documents(rag_id: str, user: dict = Depends(get_current_user)):
    """Статусы документов базы — для опроса страницей загрузки."""
    store = rag.get_store()
    card, _owner_id = _resolve(store, rag_id, user, projects.LEVEL_READ, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    return {"items": card.get("documents") or []}


@router.delete("/rag/{rag_id}/documents/{doc_id}")
def delete_document(rag_id: str, doc_id: str, user: dict = Depends(get_current_user)):
    """Удаляет документ из базы: чанки в Qdrant и запись в карточке (правка и выше).
    Саму RAG-базу этот метод не удаляет."""
    if not DOC_ID_RE.match(doc_id):
        raise HTTPException(status_code=400, detail="Некорректный идентификатор документа")
    store = rag.get_store()
    card, owner_id = _resolve(store, rag_id, user, projects.LEVEL_WRITE, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    documents = list(card.get("documents") or [])
    if not any(isinstance(d, dict) and d.get("doc_id") == doc_id for d in documents):
        raise HTTPException(status_code=404, detail="Документ не найден")
    collection_name = card.get("collection_name") or rag_id
    try:
        client = indexer.get_qdrant_client()
        if collection_name in {c.name for c in client.get_collections().collections}:
            client.delete(
                collection_name=collection_name,
                points_selector=qmodels.FilterSelector(filter=qmodels.Filter(must=[
                    qmodels.FieldCondition(key="doc_id", match=qmodels.MatchValue(value=doc_id))
                ])),
            )
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Не удалось удалить фрагменты из хранилища: {exc}"[:300])
    rest = [d for d in documents if not (isinstance(d, dict) and d.get("doc_id") == doc_id)]
    try:
        store.save(owner_id, {**card, "documents": rest})
    except RagError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    return {"ok": True, "doc_id": doc_id}


@router.patch("/rag/{rag_id}/documents/{doc_id}")
def move_document(rag_id: str, doc_id: str, payload: dict = Body(...), user: dict = Depends(get_current_user)):
    """Переносит документ в другую папку (только поле folder, чанки не трогает)."""
    if not DOC_ID_RE.match(doc_id):
        raise HTTPException(status_code=400, detail="Некорректный идентификатор документа")
    folder = payload.get("folder", "")
    if not isinstance(folder, str):
        raise HTTPException(status_code=400, detail="Поле folder ожидается строкой")
    parts = [x.strip() for x in folder.replace("\\", "/").split("/")]
    folder = "/".join(x for x in parts if x and x not in (".", ".."))[:150]
    store = rag.get_store()
    card, owner_id = _resolve(store, rag_id, user, projects.LEVEL_WRITE, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    if not any(isinstance(d, dict) and d.get("doc_id") == doc_id for d in (card.get("documents") or [])):
        raise HTTPException(status_code=404, detail="Документ не найден")
    try:
        store.update_document(owner_id, rag_id, doc_id, {"folder": folder})
    except RagError as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    return {"ok": True, "doc_id": doc_id, "folder": folder}


async def _reclassify_plan(card: dict) -> dict:
    """Свежая классификация всех чанков коллекции по текущей опубликованной
    онтологии — общая база для ontology-check (просмотр) и reclassify
    (применение). Не трогает Qdrant на запись, только читает (scroll)."""
    ontology_id = card.get("ontology_id")
    if not ontology_id:
        raise RagError("У этой RAG-базы не указана онтология")
    graph = ontologies.get_store().get_published(ontology_id)
    if not graph:
        raise RagError("Онтология не опубликована — сначала опубликуйте её")

    collection_name = card.get("collection_name") or card["id"]
    client = indexer.get_qdrant_client()
    points = indexer.fetch_collection_points(client, collection_name)

    texts = [(p.payload or {}).get("text", "") for p in points]
    vectors = [p.vector for p in points]
    old_ids = [list((p.payload or {}).get("class_ids") or []) for p in points]
    new_ids = await ontologies.classify_chunks(ontology_id, texts, vectors) if points else []

    changed_idx = [i for i, (o, n) in enumerate(zip(old_ids, new_ids)) if set(o) != set(n)]

    return {
        "graph": graph,
        "points": points,
        "old_ids": old_ids,
        "new_ids": new_ids,
        "changed_idx": changed_idx,
    }


@router.get("/rag/{rag_id}/ontology-check")
async def ontology_check(rag_id: str, user: dict = Depends(get_current_user)):
    """Сверяет RAG-базу с текущей опубликованной версией привязанной
    онтологии: что изменилось в составе классов (добавлено/удалено/
    изменено) и сколько чанков затронет актуализация — без записи."""
    store = rag.get_store()
    card, _owner_id = _resolve(store, rag_id, user, projects.LEVEL_READ, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    if not card.get("ontology_id"):
        return {"has_ontology": False}

    try:
        plan = await _reclassify_plan(card)
    except RagError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    graph = plan["graph"]
    never_classified = card.get("class_snapshot") is None
    total = len(plan["points"])
    will_change = len(plan["changed_idx"])

    if never_classified:
        # Без сохранённого снимка предыдущей версии сравнивать состав классов
        # не с чем — «добавлено всё» было бы вводящей в заблуждение подменой
        # факта «мы просто ещё не запоминали версию». Ориентируемся на то,
        # что реально показывает пересчёт: сойдутся ли новые class_ids со
        # старыми в самих чанках.
        diff = {"added": [], "removed": [], "changed": []}
        up_to_date = will_change == 0
    else:
        new_snapshot = ontologies.classification_snapshot(graph)
        old_snapshot = (card.get("class_snapshot") or {}).get("classes")
        diff = ontologies.diff_classification_snapshots(old_snapshot, new_snapshot)
        up_to_date = (
            not diff["added"] and not diff["removed"] and not diff["changed"]
            and will_change == 0
        )

    return {
        "has_ontology": True,
        "ontology_id": graph.get("id"),
        "ontology_name": graph.get("name") or graph.get("id"),
        "rag_name": card.get("name") or rag_id,
        "never_classified": never_classified,
        "up_to_date": up_to_date,
        "classes_added": diff["added"],
        "classes_removed": diff["removed"],
        "classes_changed": diff["changed"],
        "total_chunks": total,
        "will_change": will_change,
        "unchanged": total - will_change,
    }


@router.post("/rag/{rag_id}/reclassify")
async def reclassify_rag(rag_id: str, user: dict = Depends(get_current_user)):
    """Актуализирует class_ids чанков RAG-базы под текущую опубликованную
    версию привязанной онтологии (правка и выше). Текст и векторы чанков
    не трогаются — только метаданные payload. Старые значения сохраняются
    в бэкап на диске перед записью (формат общий с
    scripts/rollback_class_ids.py)."""
    store = rag.get_store()
    card, owner_id = _resolve(store, rag_id, user, projects.LEVEL_WRITE, team_kind="rag")
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    if not card.get("ontology_id"):
        raise HTTPException(status_code=400, detail="У этой RAG-базы не указана онтология")

    try:
        plan = await _reclassify_plan(card)
    except RagError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    points = plan["points"]
    old_ids = plan["old_ids"]
    new_ids = plan["new_ids"]
    changed_idx = plan["changed_idx"]
    total = len(points)

    if changed_idx:
        collection_name = card.get("collection_name") or card["id"]
        client = indexer.get_qdrant_client()

        backup_dir = Path.home() / "aios-backups" / "rag_class_ids_backfill"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_file = backup_dir / f"{rag_id}.{int(time.time())}.json"
        backup_file.write_text(
            json.dumps(
                {
                    "rag_id": rag_id,
                    "collection_name": collection_name,
                    "items": [
                        {"point_id": str(points[i].id), "class_ids": old_ids[i]}
                        for i in changed_idx
                    ],
                },
                ensure_ascii=False, indent=2,
            ),
            encoding="utf-8",
        )

        groups: dict = {}
        for i in changed_idx:
            key = tuple(sorted(new_ids[i]))
            groups.setdefault(key, []).append(points[i].id)
        for key, point_ids in groups.items():
            client.set_payload(
                collection_name=collection_name,
                payload={"class_ids": list(key)},
                points=point_ids,
            )

    # Снимок версии онтологии фиксируем в любом случае, а не только при
    # реальном изменении payload — иначе повторная проверка снова считала
    # бы базу «никогда не размеченной» (never_classified), хотя пользователь
    # уже подтвердил актуальность нажатием «Актуализировать».
    graph = plan["graph"]
    rag.get_store().set_class_snapshot(
        owner_id, rag_id,
        {
            "ontology_id": graph.get("id"),
            "ontology_updated_at": graph.get("updated_at"),
            "classes": ontologies.classification_snapshot(graph),
            "classified_at": time.time(),
        },
    )

    return {"updated": len(changed_idx), "total": total, "unchanged": total - len(changed_idx)}


@router.delete("/rag/{rag_id}")
def delete_rag(rag_id: str, _: dict = Depends(require_admin)):
    """Удаление RAG-базы — только роль «Администратор», любой по id.

    Не удаляет коллекцию в Qdrant (векторное хранилище вне зоны
    ответственности RagStore) — это делается отдельно при необходимости.
    """
    store = rag.get_store()
    card = store.find(rag_id)
    if card is None:
        raise HTTPException(status_code=404, detail="RAG-база не найдена")
    store.delete(card["user_id"], rag_id)
    return {"ok": True, "deleted": rag_id}

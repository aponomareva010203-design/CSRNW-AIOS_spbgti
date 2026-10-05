"""Инженерные возможности агентов (api/runtime/engineering_tools.py).

Отдельный модуль, а не дописывание в tools.py: tools.py уже 75 КБ, и опыт
прошлой сессии (задвоение format_table при вставке через одинаковый якорь)
говорит, что новые инструменты безопаснее заводить файлом. Подключение —
одна строка в build_default_registry:

    from .engineering_tools import engineering_tools
    for tool in engineering_tools():
        registry.register(tool)

Состав и зачем каждый нужен (обоснование приоритета — 03_БИБЛИОТЕКА):

    run_calc          расчёт в песочнице — превращает агента в инженера
    read_data_file    xlsx/csv участника → данные для расчёта
    literature_search внешние базы публикаций: OpenAlex, arXiv, PubMed
                      (через Europe PMC), Crossref, Semantic Scholar —
                      все без ключа и подписки
    compound_lookup   PubChem: свойства вещества по названию/формуле/SMILES
    make_document     docx/xlsx на выход — участник уносит файл
    notify            письмо владельцу — включает фоновых агентов

Все внешние запросы идут через httpx с коротким таймаутом и вежливым
User-Agent с почтой института (требование OpenAlex/PubChem/Crossref к
«polite pool» — даёт стабильно более высокий приоритет обслуживания).
"""

from __future__ import annotations

import asyncio
import base64
import csv
import html
import io
import json
import re
import xml.etree.ElementTree as ET
from itertools import zip_longest
from typing import Any
from urllib.parse import quote

import httpx

from . import compute_sandbox, recipes, uploads
from .config import get_settings
from .tools import Tool

HTTP_TIMEOUT = 20.0
MAX_ROWS_PREVIEW = 30
MAX_CELL_CHARS = 200
CONTACT = "aios@technolog.edu.ru"  # ← подставить реальный адрес института
USER_AGENT = f"AIOS-TechnologicalInstitute/1.0 (mailto:{CONTACT})"
ARXIV_API = "https://export.arxiv.org/api/query"
EUROPEPMC_API = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
CROSSREF_API = "https://api.crossref.org/works"
SEMANTICSCHOLAR_API = "https://api.semanticscholar.org/graph/v1/paper/search"


def _session_user(ctx: dict | None) -> str:
    return ((ctx or {}).get("session") or {}).get("user_id") or ""


def _num(value: Any) -> float | None:
    """Число из ячейки с запятой-разделителем и пробелами-разрядами."""
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# 1. run_calc — расчёт в песочнице


async def run_calc(args: dict, ctx: dict | None) -> dict:
    """Исполняет расчётный скрипт и возвращает результат вместе с кодом."""
    code = str(args.get("code") or "")
    files: dict[str, bytes] = {}

    # Данные для расчёта: вложения участника, подготовленные read_data_file.
    user_id = _session_user(ctx)
    for att_id in (args.get("data_files") or [])[:3]:
        if not user_id:
            continue
        store = uploads.get_store()
        meta = store.get(user_id, str(att_id))
        if meta is None:
            return {"ok": False, "error": f"Файл данных не найден: {att_id}"}
        path = store.content_path(user_id, meta)
        if not path.is_file():
            continue
        files[meta.get("filename") or f"{att_id}.bin"] = path.read_bytes()

    result = await compute_sandbox.run_code(code, files=files)

    # Файлы, созданные расчётом (график matplotlib, выгрузка CSV), кладём во
    # вложения владельца и отдаём модели ссылку, а не base64: содержимое
    # картинки в контексте модели бесполезно и дорого, а по ссылке участник
    # её скачает и агент покажет компонентом file_card.
    artifacts = result.pop("artifacts", []) or []
    saved = []
    for item in artifacts:
        name = str(item.get("name") or "результат")
        if not user_id:
            saved.append({"файл": name, "сохранён": False,
                          "причина": "расчёт без авторизации — файл не сохранён"})
            continue
        try:
            blob = base64.b64decode(item.get("base64") or "")
            meta = uploads.get_store().save(user_id, name, blob)
        except Exception as exc:  # noqa: BLE001 — неудача с файлом не роняет расчёт
            saved.append({"файл": name, "сохранён": False, "причина": str(exc)[:200]})
            continue
        saved.append({
            "файл": name,
            "сохранён": True,
            "размер_кб": round(len(blob) / 1024, 1),
            "ссылка": f"/api/chat/attachments/{meta['id']}",
        })
    if saved:
        result["созданные_файлы"] = saved
        result["подсказка_по_файлам"] = (
            "Покажите каждый файл человеку компонентом file_card — готовый "
            "блок, подставьте поля из созданные_файлы:\n"
            '```aios-view\n{"component": "file_card", "filename": "<файл>", '
            '"url": "<ссылка>", "size": "<размер_кб> КБ"}\n```\n'
            "Не вставляйте содержимое файла в ответ и не оставляйте ссылку "
            "голым текстом или в код-блоке: без file_card участник не увидит "
            "картинку в чате."
        )
    return result


# ---------------------------------------------------------------------------
# 2. read_data_file — таблица участника в понятном модели виде


async def read_data_file(args: dict, ctx: dict | None) -> dict:
    """Разбирает xlsx/csv из вложения: колонки, типы, статистика, образец строк.

    Модель получает не сам файл (он может быть на 5000 строк), а его паспорт:
    какие колонки, какие единицы, где числа, минимум/максимум/среднее и первые
    строки. Полные данные уходят в расчёт через run_calc(data_files=[...]).
    """
    user_id = _session_user(ctx)
    if not user_id:
        return {"ok": False, "error": "Нужен вход в AIOS: файлы доступны только владельцу."}

    att_id = str(args.get("attachment_id") or "").strip()
    store = uploads.get_store()
    meta = store.get(user_id, att_id)
    if meta is None:
        return {"ok": False, "error": f"Вложение не найдено: {att_id}. "
                                      "Попросите участника загрузить файл кнопкой 📎."}
    path = store.content_path(user_id, meta)
    if not path.is_file():
        return {"ok": False, "error": "Содержимое вложения не найдено"}
    blob = path.read_bytes()
    if not blob:
        return {"ok": False, "error": "Файл пуст"}

    filename = meta.get("filename") or att_id
    lower = filename.lower()
    sheets: dict[str, list[list]] = {}

    if lower.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"ok": False, "error": "На сервере не установлен openpyxl — "
                                          "чтение Excel недоступно."}
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        for ws in wb.worksheets[:5]:
            rows = [list(r) for r in ws.iter_rows(max_row=5000, values_only=True)]
            if rows:
                sheets[ws.title] = rows
        wb.close()
    elif lower.endswith((".csv", ".tsv", ".txt")):
        text = blob.decode("utf-8", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        sheets["данные"] = [row for row in csv.reader(io.StringIO(text), dialect)][:5000]
    else:
        return {"ok": False, "error": f"Формат не поддерживается: {filename}. "
                                      "Нужен .xlsx, .csv или .tsv."}

    out_sheets = []
    for name, rows in sheets.items():
        if not rows:
            continue
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        body = rows[1:]
        columns = []
        for idx, title in enumerate(header):
            values = [r[idx] if idx < len(r) else None for r in body]
            nums = [n for n in (_num(v) for v in values) if n is not None]
            filled = sum(1 for v in values if v not in (None, ""))
            col: dict = {"название": title or f"колонка {idx + 1}", "заполнено": filled}
            if nums and len(nums) >= max(3, filled * 0.6):
                col["тип"] = "число"
                col["мин"] = round(min(nums), 6)
                col["макс"] = round(max(nums), 6)
                col["среднее"] = round(sum(nums) / len(nums), 6)
            else:
                col["тип"] = "текст"
                seen: list[str] = []
                for v in values:
                    s = str(v).strip()
                    if s and s not in seen and len(seen) < 8:
                        seen.append(s[:MAX_CELL_CHARS])
                col["примеры"] = seen
            columns.append(col)
        out_sheets.append({
            "лист": name,
            "строк": len(body),
            "колонки": columns,
            "образец": [[str(c)[:MAX_CELL_CHARS] if c is not None else ""
                         for c in row] for row in body[:MAX_ROWS_PREVIEW]],
        })

    return {
        "ok": True,
        "файл": filename,
        "attachment_id": att_id,
        "листы": out_sheets,
        "подсказка": "Полные данные доступны в расчёте: run_calc(code=…, "
                     f"data_files=['{att_id}']) — файл окажется рядом со "
                     f"скриптом под именем «{filename}».",
    }


# ---------------------------------------------------------------------------
# 3. literature_search — OpenAlex, arXiv, PubMed (Europe PMC), Crossref,
#    Semantic Scholar


def _clean_text(text: Any) -> str:
    """Текст без тегов: Europe PMC и Crossref кладут разметку прямо в
    аннотации (<sub>, <jats:p>), arXiv — переносы строк внутри заголовка.
    Инлайн-разметка (CeCl<sub>3</sub>) склеивается без пробела, блочная
    (<jats:p>) — разделяется."""
    text = str(text or "")
    text = re.sub(r"</?(?:jats:)?(?:p|sec|h1|h2|h3|title|br)\b[^>]*/?>", " ", text,
                  flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    return re.sub(r"\s+", " ", html.unescape(text)).strip()


def _openalex_works(data: dict) -> list[dict]:
    works = []
    for w in data.get("results", []):
        primary = w.get("primary_location") or {}
        source = (primary.get("source") or {}).get("display_name") or ""
        authors = [
            (a.get("author") or {}).get("display_name", "")
            for a in (w.get("authorships") or [])[:5]
        ]
        # abstract_inverted_index → обычный текст
        abstract = ""
        inv = w.get("abstract_inverted_index")
        if isinstance(inv, dict):
            positions: dict[int, str] = {}
            for word, places in inv.items():
                for p in places:
                    positions[p] = word
            abstract = " ".join(positions[k] for k in sorted(positions))[:1200]
        works.append({
            "название": w.get("display_name"),
            "год": w.get("publication_year"),
            "авторы": [a for a in authors if a],
            "журнал": source,
            "цитирований": w.get("cited_by_count"),
            "doi": (w.get("doi") or "").replace("https://doi.org/", ""),
            "открытый_доступ": bool((w.get("open_access") or {}).get("is_oa")),
            "pdf": (primary.get("pdf_url") or ""),
            "ссылка": w.get("doi") or w.get("id"),
            "аннотация": abstract,
        })
    return works


async def _search_openalex(args: dict, query: str, limit: int) -> dict:
    params: dict[str, Any] = {"search": query, "per-page": limit, "mailto": CONTACT}
    filters = []
    year_from = args.get("year_from")
    if year_from:
        filters.append(f"from_publication_date:{int(year_from)}-01-01")
    if args.get("open_access_only"):
        filters.append("is_oa:true")
    import datetime as _dt
    filters.append("to_publication_date:" + _dt.date.today().isoformat())  # без записей с датой из будущего
    if filters:
        params["filter"] = ",".join(filters)
    if args.get("sort") == "citations":
        params["sort"] = "cited_by_count:desc"
    elif args.get("sort") == "newest":
        params["sort"] = "publication_date:desc"

    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        import asyncio as _asyncio
        resp = None
        for attempt in range(3):  # 429 и 5xx: повтор с паузой 1,5 и 3 с
            resp = await client.get("https://api.openalex.org/works", params=params)
            if resp.status_code not in (429, 500, 502, 503, 504):
                break
            if attempt < 2:
                await _asyncio.sleep(1.5 * (attempt + 1))
        if resp.status_code != 200:
            return {"ok": False, "error": f"OpenAlex ответил {resp.status_code} после 3 попыток"}
        data = resp.json()

    return {
        "ok": True,
        "запрос": query,
        "найдено_всего": (data.get("meta") or {}).get("count", 0),
        "работы": _openalex_works(data),
        "источник": "OpenAlex",
    }


def _arxiv_query(query: str) -> str:
    """arXiv ждёт поисковый синтаксис: термины через AND, каждое с префиксом
    поля. Фразы в кавычках и склейки без поля на выдаче дают ноль."""
    terms = [t for t in query.split() if t]
    return " AND ".join(f"all:{t}" for t in terms) or query


def _arxiv_works(xml_text: str) -> tuple[list[dict], int]:
    """Разбор Atom-выдачи arXiv → (работы, всего найдено)."""
    atom = "{http://www.w3.org/2005/Atom}"
    ax = "{http://arxiv.org/schemas/atom}"
    opensearch = "{http://a9.com/-/spec/opensearch/1.1/}"
    root = ET.fromstring(xml_text)
    total = int(root.findtext(opensearch + "totalResults") or 0)
    works = []
    for entry in root.findall(atom + "entry"):
        published = entry.findtext(atom + "published") or ""
        year = int(published[:4]) if published[:4].isdigit() else None
        pdf = ""
        for link in entry.findall(atom + "link"):
            if link.get("type") == "application/pdf":
                pdf = link.get("href") or ""
                break
        works.append({
            "название": _clean_text(entry.findtext(atom + "title")),
            "год": year,
            "авторы": [a.findtext(atom + "name") or ""
                       for a in entry.findall(atom + "author")][:5],
            "журнал": _clean_text(entry.findtext(ax + "journal_ref")) or "arXiv (препринт)",
            "цитирований": None,  # данных о цитированиях у arXiv нет
            "doi": entry.findtext(ax + "doi") or "",
            "открытый_доступ": True,  # препринты arXiv всегда открыты
            "pdf": pdf,
            "ссылка": entry.findtext(atom + "id") or "",
            "аннотация": _clean_text(entry.findtext(atom + "summary"))[:1200],
        })
    return works, total


async def _search_arxiv(args: dict, query: str, limit: int) -> dict:
    params: dict[str, Any] = {"search_query": _arxiv_query(query),
                              "start": 0, "max_results": limit,
                              "sortBy": "relevance"}
    if args.get("sort") == "newest":
        params["sortBy"], params["sortOrder"] = "submittedDate", "descending"
    # sort=citations у arXiv не существует — молча ищем по релевантности

    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get(ARXIV_API, params=params, follow_redirects=True)
        if resp.status_code != 200:
            return {"ok": False, "error": f"arXiv ответил {resp.status_code}"}
    try:
        works, total = _arxiv_works(resp.text)
    except ET.ParseError:
        return {"ok": False, "error": "arXiv вернул нечитаемый ответ"}

    year_from = args.get("year_from")
    if year_from:
        works = [w for w in works if (w.get("год") or 0) >= int(year_from)]

    return {"ok": True, "запрос": query, "найдено_всего": total,
            "работы": works, "источник": "arXiv"}


def _europepmc_works(data: dict) -> list[dict]:
    works = []
    for r in ((data.get("resultList") or {}).get("result") or []):
        pdf = ""
        for url in ((r.get("fullTextUrlList") or {}).get("fullTextUrl") or []):
            if url.get("documentStyle") == "pdf":
                pdf = url.get("url") or ""
                if url.get("availabilityCode") == "OA":
                    break  # предпочитаем PDF открытого доступа
        doi = r.get("doi") or ""
        year = r.get("pubYear") or (r.get("firstPublicationDate") or "")[:4]
        journal = ((r.get("journalInfo") or {}).get("journal") or {}).get("title") or ""
        works.append({
            "название": _clean_text(r.get("title")),
            "год": int(year) if str(year or "").isdigit() else None,
            "авторы": [a.strip() for a in (r.get("authorString") or "").split(",")
                       if a.strip()][:5],
            "журнал": _clean_text(journal),
            "цитирований": r.get("citedByCount"),
            "doi": doi,
            "открытый_доступ": r.get("isOpenAccess") == "Y",
            "pdf": pdf,
            "ссылка": (f"https://doi.org/{doi}" if doi else
                       f"https://europepmc.org/article/{r.get('source')}/{r.get('id')}"),
            "аннотация": _clean_text(r.get("abstractText"))[:1200],
        })
    return works


async def _search_pubmed(args: dict, query: str, limit: int) -> dict:
    """PubMed через Europe PMC: у E-utilities NCBI лимиты скромнее и ответ
    XML, а Europe PMC отдаёт те же записи MEDLINE открытым JSON без ключа."""
    q = query
    if args.get("open_access_only"):
        q += " AND OPEN_ACCESS:y"
    params: dict[str, Any] = {"query": q, "format": "json", "pageSize": limit,
                              "resultType": "core"}
    if args.get("sort") == "citations":
        params["sort"] = "CITED desc"
    elif args.get("sort") == "newest":
        params["sort"] = "P_PDATE_D desc"

    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get(EUROPEPMC_API, params=params)
        if resp.status_code != 200:
            return {"ok": False, "error": f"PubMed (Europe PMC) ответил {resp.status_code}"}
        data = resp.json()

    return {"ok": True, "запрос": query, "найдено_всего": data.get("hitCount", 0),
            "работы": _europepmc_works(data), "источник": "PubMed (Europe PMC)"}


def _crossref_works(data: dict) -> list[dict]:
    works = []
    for it in ((data.get("message") or {}).get("items") or []):
        doi = it.get("DOI") or ""
        pdf = ""
        for link in it.get("link") or []:
            if "pdf" in (link.get("content-type") or ""):
                pdf = link.get("URL") or ""
                break
        date = (it.get("published") or it.get("issued") or {}).get("date-parts") or []
        year = date[0][0] if date and date[0] else None
        licenses = json.dumps(it.get("license") or [], ensure_ascii=False).lower()
        authors = [f"{a.get('given') or ''} {a.get('family') or ''}".strip()
                   for a in (it.get("author") or [])[:5]]
        works.append({
            "название": _clean_text((it.get("title") or [""])[0]),
            "год": year,
            "авторы": [a for a in authors if a],
            "журнал": _clean_text((it.get("container-title") or [""])[0]),
            "цитирований": it.get("is-referenced-by-count"),
            "doi": doi,
            "открытый_доступ": "creativecommons" in licenses,
            "pdf": pdf,
            "ссылка": f"https://doi.org/{doi}" if doi else (it.get("URL") or ""),
            "аннотация": _clean_text(it.get("abstract"))[:1200],
        })
    return works


async def _search_crossref(args: dict, query: str, limit: int) -> dict:
    params: dict[str, Any] = {
        "query": query, "rows": limit, "mailto": CONTACT,
        "select": "title,author,container-title,published,issued,"
                  "is-referenced-by-count,DOI,abstract,link,URL,license",
    }
    if args.get("sort") == "citations":
        params["sort"], params["order"] = "is-referenced-by-count", "desc"
    elif args.get("sort") == "newest":
        params["sort"], params["order"] = "published", "desc"
    year_from = args.get("year_from")
    if year_from:
        params["filter"] = f"from-pub-date:{int(year_from)}-01-01"

    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get(CROSSREF_API, params=params)
        if resp.status_code != 200:
            return {"ok": False, "error": f"Crossref ответил {resp.status_code}"}
        data = resp.json()

    message = data.get("message") or {}
    return {"ok": True, "запрос": query,
            "найдено_всего": message.get("total-results", 0),
            "работы": _crossref_works(data), "источник": "Crossref"}


_SEMANTICSCHOLAR_FIELDS = ("title,abstract,tldr,year,venue,citationCount,"
                           "isOpenAccess,openAccessPdf,externalIds,authors,url")


def _semanticscholar_works(data: dict) -> list[dict]:
    works = []
    for w in data.get("data") or []:
        doi = ((w.get("externalIds") or {}).get("DOI")) or ""
        # abstract бывает null — тогда берём TLDR-выжимку, она почти всегда есть
        abstract = _clean_text(w.get("abstract"))[:1200] or \
            _clean_text((w.get("tldr") or {}).get("text"))[:1200]
        works.append({
            "название": _clean_text(w.get("title")),
            "год": w.get("year"),
            "авторы": [a.get("name") or "" for a in (w.get("authors") or [])][:5],
            "журнал": _clean_text(w.get("venue")),
            "цитирований": w.get("citationCount"),
            "doi": doi,
            "открытый_доступ": bool(w.get("isOpenAccess")),
            "pdf": (w.get("openAccessPdf") or {}).get("url") or "",
            "ссылка": f"https://doi.org/{doi}" if doi else (w.get("url") or ""),
            "аннотация": abstract,
        })
    return works


async def _search_semanticscholar(args: dict, query: str, limit: int) -> dict:
    """Semantic Scholar Graph API: без ключа лимит частоты запросов жёсткий,
    403/429 честно называем и советуем, что делать."""
    params: dict[str, Any] = {"query": query, "limit": limit,
                              "fields": _SEMANTICSCHOLAR_FIELDS}
    year_from = args.get("year_from")
    if year_from:
        params["year"] = f"{int(year_from)}-"
    if args.get("open_access_only"):
        params["openAccessPdf"] = "any"
    if args.get("sort") == "citations":
        params["sort"] = "citationCount:desc"
    elif args.get("sort") == "newest":
        params["sort"] = "publicationDate:desc"

    headers = {"User-Agent": USER_AGENT}
    api_key = get_settings().semanticscholar_api_key
    if api_key:
        headers["x-api-key"] = api_key

    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers=headers) as client:
        resp = await client.get(SEMANTICSCHOLAR_API, params=params)
        if resp.status_code == 403:
            return {"ok": False, "error": "Semantic Scholar ответил 403 — сервис "
                    "не пускает этот сервер без ключа: оформите бесплатный ключ "
                    "(semanticscholar.org/product/api) и задайте его в "
                    "AIOS_SEMANTICSCHOLAR_API_KEY, либо возьмите другую базу"}
        if resp.status_code == 429:
            return {"ok": False, "error": "Semantic Scholar ответил 429 — лимит "
                    "запросов исчерпан: повторите позже или возьмите другую базу"}
        if resp.status_code != 200:
            return {"ok": False, "error": f"Semantic Scholar ответил {resp.status_code}"}
        data = resp.json()

    return {"ok": True, "запрос": query,
            "найдено_всего": data.get("total") or 0,
            "работы": _semanticscholar_works(data), "источник": "Semantic Scholar"}


# ---- дополнительные источники (без ключей; вызываются только явно, в source=all не входят) ----

def _strip_html(text) -> str:
    import re as _re
    return _clean_text(_re.sub(r"<[^>]+>", " ", str(text or "")))


def _year_int(value):
    value = str(value or "")[:4]
    return int(value) if value.isdigit() else None


def _after_year(works: list, args: dict) -> list:
    """year_from для источников без серверного фильтра — по выдаче."""
    try:
        year_from = int(args.get("year_from") or 0)
    except (TypeError, ValueError):
        return works
    if not year_from:
        return works
    return [w for w in works if not w.get("год") or w["год"] >= year_from]


async def _search_zenodo(args: dict, query: str, limit: int) -> dict:
    params: dict[str, Any] = {"q": query, "size": limit, "sort": "bestmatch"}
    if args.get("sort") == "newest":
        params["sort"] = "mostrecent"
    if args.get("open_access_only"):
        params["access_right"] = "open"
    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get("https://zenodo.org/api/records", params=params)
        if resp.status_code != 200:
            return {"ok": False, "error": f"Zenodo ответил {resp.status_code}"}
        data = resp.json()
    hits = data.get("hits") or {}
    total = hits.get("total") or 0
    if isinstance(total, dict):
        total = total.get("value", 0)
    works = []
    for h in hits.get("hits") or []:
        md = h.get("metadata") or {}
        doi = h.get("doi") or md.get("doi") or ""
        pdf = ""
        for f in h.get("files") or []:
            if str(f.get("key") or "").lower().endswith(".pdf"):
                pdf = (f.get("links") or {}).get("self") or ""
                break
        kind = (md.get("resource_type") or {}).get("title") or "запись"
        works.append({
            "название": _clean_text(md.get("title")),
            "год": _year_int(md.get("publication_date")),
            "авторы": [a.get("name") for a in (md.get("creators") or [])[:5] if a.get("name")],
            "журнал": f"Zenodo ({kind})",
            "цитирований": None,
            "doi": doi,
            "открытый_доступ": md.get("access_right") == "open",
            "pdf": pdf,
            "ссылка": (h.get("links") or {}).get("html") or (f"https://doi.org/{doi}" if doi else ""),
            "аннотация": _strip_html(md.get("description"))[:1200],
        })
    return {"ok": True, "запрос": query, "найдено_всего": total,
            "работы": _after_year(works, args), "источник": "Zenodo"}


async def _search_doaj(args: dict, query: str, limit: int) -> dict:
    from urllib.parse import quote
    params: dict[str, Any] = {"pageSize": limit}
    if args.get("sort") == "newest":
        params["sort"] = "created_date:desc"
    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get("https://doaj.org/api/search/articles/" + quote(query, safe=""),
                                params=params)
        if resp.status_code != 200:
            return {"ok": False, "error": f"DOAJ ответил {resp.status_code}"}
        data = resp.json()
    works = []
    for r in data.get("results") or []:
        b = r.get("bibjson") or {}
        doi = next((i.get("id") for i in (b.get("identifier") or [])
                    if str(i.get("type")).lower() == "doi"), "") or ""
        pdf = next((l.get("url") for l in (b.get("link") or [])
                    if str(l.get("type")).lower() == "fulltext"), "") or ""
        works.append({
            "название": _clean_text(b.get("title")),
            "год": _year_int(b.get("year")),
            "авторы": [a.get("name") for a in (b.get("author") or [])[:5] if a.get("name")],
            "журнал": _clean_text((b.get("journal") or {}).get("title")),
            "цитирований": None,
            "doi": doi,
            "открытый_доступ": True,  # DOAJ индексирует только открытые журналы
            "pdf": pdf,
            "ссылка": f"https://doi.org/{doi}" if doi else pdf,
            "аннотация": _clean_text(b.get("abstract"))[:1200],
        })
    return {"ok": True, "запрос": query, "найдено_всего": data.get("total") or 0,
            "работы": _after_year(works, args), "источник": "DOAJ"}


async def _search_datacite(args: dict, query: str, limit: int) -> dict:
    params: dict[str, Any] = {"query": query, "page[size]": limit}
    if args.get("sort") == "newest":
        params["sort"] = "-created"
    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get("https://api.datacite.org/dois", params=params)
        if resp.status_code != 200:
            return {"ok": False, "error": f"DataCite ответил {resp.status_code}"}
        data = resp.json()
    works = []
    for item in data.get("data") or []:
        a = item.get("attributes") or {}
        doi = a.get("doi") or ""
        kind = (a.get("types") or {}).get("resourceTypeGeneral") or "запись"
        works.append({
            "название": _clean_text(((a.get("titles") or [{}])[0]).get("title")),
            "год": _year_int(a.get("publicationYear")),
            "авторы": [c.get("name") for c in (a.get("creators") or [])[:5] if c.get("name")],
            "журнал": f"{a.get('publisher') or 'DataCite'} ({kind})",
            "цитирований": a.get("citationCount"),
            "doi": doi,
            "открытый_доступ": None,
            "pdf": "",
            "ссылка": f"https://doi.org/{doi}" if doi else (a.get("url") or ""),
            "аннотация": _strip_html(((a.get("descriptions") or [{}])[0]).get("description"))[:1200],
        })
    return {"ok": True, "запрос": query,
            "найдено_всего": (data.get("meta") or {}).get("total") or 0,
            "работы": _after_year(works, args), "источник": "DataCite"}


_EXPLICIT_ONLY_SOURCES = ("zenodo", "doaj", "datacite")  # в source=all не входят


_LITERATURE_SOURCES = {
    "openalex": _search_openalex,
    "arxiv": _search_arxiv,
    "pubmed": _search_pubmed,
    "crossref": _search_crossref,
    "semanticscholar": _search_semanticscholar,
    "zenodo": _search_zenodo,
    "doaj": _search_doaj,
    "datacite": _search_datacite,
}
_SOURCE_LABELS = {"openalex": "OpenAlex", "arxiv": "arXiv",
                  "pubmed": "PubMed (Europe PMC)", "crossref": "Crossref",
                  "semanticscholar": "Semantic Scholar", "zenodo": "Zenodo",
                  "doaj": "DOAJ", "datacite": "DataCite"}
ALL_MODE_MAX = 20  # потолок сводного списка, чтобы не съесть контекст модели


async def literature_search(args: dict, ctx: dict | None) -> dict:
    """Поиск научных работ во внешних базах: OpenAlex, arXiv, PubMed (через
    Europe PMC), Crossref, Semantic Scholar.

    Все открыты, без ключей и подписок — в отличие от Scopus/Web of Science.
    Отдаёт то, чего нет в RAG-базе участника: мировой контекст его задачи.
    По умолчанию — OpenAlex; source="all" сводит все пять баз в один
    список с чередованием источников.
    """
    query = str(args.get("query") or "").strip()
    if not query:
        return {"ok": False, "error": "Пустой поисковый запрос"}
    raw_limit = args.get("limit")
    try:
        limit = int(raw_limit) if raw_limit else 10
    except (TypeError, ValueError):
        limit = 10
    limit = max(1, min(limit, 25))

    source = str(args.get("source") or "openalex").strip().lower()
    if source == "europepmc":
        source = "pubmed"

    if source == "all":
        names = [n for n in _LITERATURE_SOURCES if n not in _EXPLICIT_ONLY_SOURCES]
        results = await asyncio.gather(
            *(_LITERATURE_SOURCES[n](args, query, limit) for n in names),
            return_exceptions=True)
        by_source: dict[str, dict] = {}
        groups: list[list[dict]] = []
        for name, res in zip(names, results):
            label = _SOURCE_LABELS[name]
            if isinstance(res, BaseException):
                by_source[label] = {"ошибка": str(res)[:200] or type(res).__name__}
            elif not res.get("ok"):
                by_source[label] = {"ошибка": str(res.get("error") or "нет ответа")[:200]}
            else:
                by_source[label] = {"найдено_всего": res.get("найдено_всего"),
                                    "в_выдаче": len(res.get("работы") or [])}
                for work in res["работы"]:
                    work["источник"] = label
                groups.append(res["работы"])
        if not groups:
            return {"ok": False, "error": "Ни один источник не ответил",
                    "по_источникам": by_source}
        # Чередуем базы раунд-робином: в начале списка — работы всех источников
        merged = [w for row in zip_longest(*groups) for w in row if w is not None]
        return {
            "ok": True,
            "запрос": query,
            "источник": " + ".join(_SOURCE_LABELS[n] for n in names),
            "найдено_всего": sum((s.get("найдено_всего") or 0)
                                 for s in by_source.values()),
            "работы": merged[:ALL_MODE_MAX],
            "по_источникам": by_source,
            "подсказка": "Происхождение каждой работы — в её поле «источник»; "
                         "подробную выдачу одной базы берите отдельным вызовом.",
        }

    fn = _LITERATURE_SOURCES.get(source)
    if fn is None:
        return {"ok": False, "error": f"Неизвестный источник: {source}",
                "доступные": ["openalex", "arxiv", "pubmed", "crossref",
                              "semanticscholar", "zenodo", "doaj", "datacite",
                              "all"]}
    return await fn(args, query, limit)


# ---------------------------------------------------------------------------
# 4. compound_lookup — PubChem


async def compound_lookup(args: dict, ctx: dict | None) -> dict:
    """Свойства вещества из PubChem по названию, формуле, CAS или SMILES."""
    name = str(args.get("name") or "").strip()
    if not name:
        return {"ok": False, "error": "Не указано вещество"}

    namespace = str(args.get("namespace") or "name").strip()
    if namespace not in ("name", "smiles", "formula", "cid"):
        namespace = "name"
    props = ("MolecularFormula,MolecularWeight,CanonicalSMILES,InChIKey,"
             "IUPACName,XLogP,TPSA,HBondDonorCount,HBondAcceptorCount,"
             "RotatableBondCount,Complexity,Charge")
    base = "https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound"
    url = f"{base}/{namespace}/{quote(name, safe='')}/property/{props}/JSON"

    async with httpx.AsyncClient(timeout=HTTP_TIMEOUT,
                                 headers={"User-Agent": USER_AGENT}) as client:
        resp = await client.get(url)
        if resp.status_code == 404:
            return {"ok": False, "error": f"В PubChem не найдено: {name}"}
        if resp.status_code != 200:
            return {"ok": False, "error": f"PubChem ответил {resp.status_code}"}
        data = resp.json()

    items = ((data.get("PropertyTable") or {}).get("Properties") or [])
    if not items:
        return {"ok": False, "error": f"В PubChem не найдено: {name}"}
    p = items[0]
    return {
        "ok": True,
        "запрос": name,
        "cid": p.get("CID"),
        "название_iupac": p.get("IUPACName"),
        "формула": p.get("MolecularFormula"),
        "молярная_масса": p.get("MolecularWeight"),
        # PubChem переименовал CanonicalSMILES в ConnectivitySMILES и отдаёт
        # ответ под новым ключом даже на запрос по старому имени — читаем оба.
        "smiles": p.get("ConnectivitySMILES") or p.get("CanonicalSMILES") or p.get("SMILES"),
        "inchikey": p.get("InChIKey"),
        "logP": p.get("XLogP"),
        "TPSA": p.get("TPSA"),
        "доноры_H": p.get("HBondDonorCount"),
        "акцепторы_H": p.get("HBondAcceptorCount"),
        "вращающиеся_связи": p.get("RotatableBondCount"),
        "заряд": p.get("Charge"),
        "ссылка": f"https://pubchem.ncbi.nlm.nih.gov/compound/{p.get('CID')}",
        "источник": "PubChem",
    }


# ---------------------------------------------------------------------------
# 5. make_document — файл на выход


async def make_document(args: dict, ctx: dict | None) -> dict:
    """Собирает docx или xlsx и кладёт его во вложения пользователя.

    Участник уносит с семинара файл, а не ссылку на переписку — это, по
    наблюдению за похожими внедрениями, сильнее всего влияет на то, вернётся
    ли он на платформу.

    Формат содержимого (`blocks`) намеренно простой, чтобы модель заполняла
    его без ошибок: heading / text / table / list.
    """
    user_id = _session_user(ctx)
    if not user_id:
        return {"ok": False, "error": "Нужен вход в AIOS: документ сохраняется владельцу."}

    kind = str(args.get("format") or "docx").lower()
    title = str(args.get("title") or "Документ").strip()[:200]
    blocks = args.get("blocks") or []
    safe_title = re.sub(r"[^\w\-. ]+", "_", title, flags=re.UNICODE).strip() or "Документ"

    buffer = io.BytesIO()
    if kind == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            return {"ok": False, "error": "На сервере не установлен openpyxl"}
        wb = Workbook()
        ws = wb.active
        ws.title = safe_title[:31]
        row = 1
        for block in blocks:
            btype = (block or {}).get("type")
            if btype == "heading":
                ws.cell(row=row, column=1, value=str(block.get("text", "")))
                row += 2
            elif btype == "table":
                columns = block.get("columns") or []
                for c, head in enumerate(columns, start=1):
                    ws.cell(row=row, column=c, value=str(head))
                row += 1
                for item in block.get("rows") or []:
                    for c, head in enumerate(columns, start=1):
                        ws.cell(row=row, column=c, value=(item or {}).get(head, ""))
                    row += 1
                row += 1
            elif btype in ("text", "list"):
                for line in ([block.get("text", "")] if btype == "text" else block.get("items") or []):
                    ws.cell(row=row, column=1, value=str(line))
                    row += 1
                row += 1
        wb.save(buffer)
        filename = f"{safe_title}.xlsx"
    else:
        try:
            from docx import Document
        except ImportError:
            return {"ok": False, "error": "На сервере не установлен python-docx"}
        doc = Document()
        doc.add_heading(title, level=0)
        for block in blocks:
            btype = (block or {}).get("type")
            if btype == "heading":
                doc.add_heading(str(block.get("text", "")), level=int(block.get("level") or 1))
            elif btype == "text":
                doc.add_paragraph(str(block.get("text", "")))
            elif btype == "list":
                for item in block.get("items") or []:
                    doc.add_paragraph(str(item), style="List Bullet")
            elif btype == "table":
                columns = block.get("columns") or []
                rows = block.get("rows") or []
                if columns:
                    table = doc.add_table(rows=1, cols=len(columns))
                    table.style = "Light Grid Accent 1"
                    for c, head in enumerate(columns):
                        table.rows[0].cells[c].text = str(head)
                    for item in rows:
                        cells = table.add_row().cells
                        for c, head in enumerate(columns):
                            cells[c].text = str((item or {}).get(head, ""))
        doc.save(buffer)
        filename = f"{safe_title}.docx"

    blob = buffer.getvalue()
    # UploadStore.save проверяет расширение по ALLOWED_SUFFIXES — для .xlsx
    # его надо расширить (см. 08_ИНСТРУКЦИЯ, шаг «форматы вложений»).
    meta = uploads.get_store().save(user_id, filename, blob)
    return {
        "ok": True,
        "файл": filename,
        "размер_кб": round(len(blob) / 1024, 1),
        "attachment_id": meta.get("id"),
        "скачать": f"/api/chat/attachments/{meta.get('id')}",
        "подсказка": (
            "Покажите файл человеку компонентом file_card:\n"
            '```aios-view\n{"component": "file_card", "filename": "<файл>", '
            '"url": "<скачать>", "size": "<размер_кб> КБ"}\n```'
        ),
    }


# ---------------------------------------------------------------------------
# 6. notify — письмо владельцу (делает фонового агента настоящим)


# Лимит писем: не больше _NOTIFY_LIMIT за _NOTIFY_WINDOW секунд на пользователя.
# Счётчик в памяти процесса, при перезапуске сервиса обнуляется.
_NOTIFY_LOG: dict = {}
_NOTIFY_LIMIT = 10
_NOTIFY_WINDOW = 3600


async def notify(args: dict, ctx: dict | None) -> dict:
    """Отправляет пользователю (владельцу сессии) письмо с результатом работы.

    Без этого фоновый агент по расписанию бесполезен: он что-то делает, но
    узнать об этом можно, только зайдя на его страницу.
    """
    from . import mailer
    from .users import UserStore

    user_id = _session_user(ctx)
    if not user_id:
        return {"ok": False, "error": "Некому отправлять: нет владельца сессии"}
    user = UserStore(get_settings().runtime_dir / "users").get(user_id)
    # отдельного поля email у учётной записи нет: логин — и есть адрес
    # (LOGIN_RE пропускает только e-mail); поле оставлено для старых записей
    email = (user or {}).get("email") or (user or {}).get("login") or ""
    if not email:
        return {"ok": False, "error": "У владельца не указана почта"}

    subject = str(args.get("subject") or "Сообщение от вашего агента AIOS")[:200]
    body = str(args.get("body") or "").strip()[:8000]
    if not body:
        return {"ok": False, "error": "Пустое сообщение"}

    import os
    import time
    now = time.time()
    recent = [t for t in _NOTIFY_LOG.get(user_id, []) if now - t < _NOTIFY_WINDOW]
    _NOTIFY_LOG[user_id] = recent
    if len(recent) >= _NOTIFY_LIMIT:
        return {"ok": False, "error": (
            f"Лимит писем исчерпан: не больше {_NOTIFY_LIMIT} в час. "
            "Соберите результаты в одно письмо и повторите позже."
        )}
    session_id = ((ctx or {}).get("session") or {}).get("id") or ""
    if session_id:
        base = os.environ.get("AIOS_PUBLIC_URL", "").rstrip("/")
        link = f"{base}/chat.html?session={session_id}" if base else f"chat.html?session={session_id}"
        body += f"\n\n—\nСессия AIOS: {session_id}\n{link}"

    # Личный SMTP владельца перекрывает админский только для его писем:
    # у личного канала свои лимиты и своё «от кого» (см. user_mail.py).
    from . import mail_settings, user_mail
    from .mail_settings import MailConfig

    personal = user_mail.personal_mail(user_id)
    if personal:
        cfg = MailConfig(
            host=str(personal["host"]),
            port=int(personal.get("port") or 465),
            tls=str(personal.get("tls") or "ssl"),
            user=str(personal.get("user") or ""),
            password=str(personal.get("password") or ""),
            from_email=str(personal.get("from_email") or personal.get("user") or ""),
            from_name=str(personal.get("from_name") or "AIOS"),
            source="personal",
        )
        send = mailer.send_with_config
    else:
        cfg = mail_settings.effective_mail_config()
        send = mailer.send_bulk_html_email

    # отправка синхронная (smtplib) — уводим в поток, чтобы не блокировать
    # цикл событий на время SMTP-сессии
    try:
        outcome = await asyncio.to_thread(send, [email], subject, body, cfg)
    except mailer.MailError as exc:
        if personal is None:
            return {"ok": False, "error": (
                "Почта не настроена: письма недоступны. Настройте личную почту "
                "или попросите администратора заполнить SMTP платформы."
            )}
        return {"ok": False, "error": f"Почта не отправилась: {exc}"}
    ok = bool(outcome.get("sent"))
    if ok:
        _NOTIFY_LOG.setdefault(user_id, []).append(now)
    return {"ok": ok, "кому": email, "канал": cfg.source or "platform",
            "error": None if ok else str(outcome.get("failed"))}



# ---------------------------------------------------------------------------
# 7. get_recipe — методики института


async def get_recipe(args: dict, ctx: dict | None) -> dict:
    """Проверенный порядок расчёта из библиотеки института.

    Без имени — список того, что есть. С именем — полный текст: метод,
    рабочий код с верными сигнатурами, пороги и типовые ошибки.
    """
    name = str(args.get("name") or "").strip()
    if not name:
        return {"ok": True, "методики": recipes.index(),
                "подсказка": "Вызовите get_recipe с полем name, чтобы прочитать нужную."}
    found = recipes.get(name)
    if found is None:
        return {"ok": False,
                "error": f"Методики «{name}» нет.",
                "доступные": [r["name"] for r in recipes.index()]}
    return {"ok": True, "название": found["title"], "когда": found["when"],
            "текст": found["body"]}


# ---------------------------------------------------------------------------
# Регистрация


def engineering_tools() -> list[Tool]:
    """Инструменты для регистрации в общем реестре."""
    tools = [
        Tool(
            name="read_data_file",
            description="Читает таблицу участника (Excel .xlsx или .csv), загруженную "
            "в чат кнопкой 📎: возвращает состав колонок, их тип, диапазоны значений "
            "и первые строки. Вызывайте ПЕРЕД расчётом, чтобы узнать, что в файле; "
            "сами данные потом передавайте в run_calc через data_files. "
            "attachment_id (att-…) берите из сообщения участника, не выдумывайте.",
            parameters={
                "type": "object",
                "properties": {
                    "attachment_id": {"type": "string", "description": "id вложения (att-…)"},
                },
                "required": ["attachment_id"],
            },
            handler=read_data_file,
        ),
        Tool(
            name="literature_search",
            description="Ищет научные публикации во внешних открытых базах. Выбирайте "
            "source по теме запроса: openalex (по умолчанию) — все дисциплины, 250+ млн "
            "работ, есть цитирования — базовый выбор; arxiv — свежие препринты физики, "
            "материаловедения, CS/ML (цитирований нет, зато новее всех); pubmed — "
            "медицина, биология, биоматериалы, токсикология (через Europe PMC); crossref "
            "— метаданные DOI всех издательств, надёжен по точным названиям; "
            "semanticscholar — все дисциплины с сильным покрытием CS/ML, свежие счётчики "
            "цитирований, TLDR-выжимка вместо отсутствующей аннотации; all — "
            "сводка сразу из пяти баз (до 20 работ вперемешку, происхождение каждой "
            "— в поле «источник»); zenodo — препринты, датасеты и приложения к статьям "
            "(данные экспериментов, код); doaj — рецензируемые статьи журналов открытого "
            "доступа; datacite — датасеты и программы с DOI; zenodo, doaj и datacite "
            "в source=all не входят — вызывайте их отдельно. Возвращает название, авторов, журнал, год, "
            "цитирования, DOI, ссылку на PDF в открытом доступе, аннотацию. Используйте, "
            "когда нужны данные, которых НЕТ в базе знаний проекта: мировой уровень "
            "показателя, чужие условия синтеза, аналоги, обзор по теме. Классика по "
            "теме — sort=citations, свежие публикации (мониторинг) — sort=newest. "
            "Запросы на английском — по-русски в мировой литературе почти ничего нет. "
            "Оформляйте результат компонентом source_list или таблицей, всегда с DOI.",
            parameters={
                "type": "object",
                "properties": {
                    "query": {"type": "string", "description": "запрос по-английски — "
                              "в мировой литературе русские запросы почти ничего не находят"},
                    "source": {"type": "string",
                               "enum": ["openalex", "arxiv", "pubmed", "crossref",
                                        "semanticscholar", "zenodo", "doaj",
                                        "datacite", "all"],
                               "description": "какая база; по умолчанию openalex"},
                    "limit": {"type": "integer", "description": "сколько работ (1-25), по умолчанию 10"},
                    "year_from": {"type": "integer", "description": "не раньше этого года "
                                  "(в OpenAlex, Crossref и Semantic Scholar — фильтром на "
                                  "сервере, в arXiv — по выдаче, в PubMed игнорируется)"},
                    "open_access_only": {"type": "boolean", "description": "только доступные полнотекстово"},
                    "sort": {"type": "string", "enum": ["relevance", "citations", "newest"],
                             "description": "citations — сначала самые цитируемые, "
                             "newest — сначала свежие"},
                },
                "required": ["query"],
            },
            handler=literature_search,
        ),
        Tool(
            name="compound_lookup",
            description="Справка по химическому веществу из PubChem: формула, молярная "
            "масса, SMILES, InChIKey, logP, TPSA, доноры/акцепторы водородной связи, "
            "число вращающихся связей. Используйте, когда нужны точные паспортные "
            "свойства реагента или продукта — не берите их из головы.",
            parameters={
                "type": "object",
                "properties": {
                    "name": {"type": "string", "description": "название, формула, CAS, SMILES или CID"},
                    "namespace": {"type": "string", "enum": ["name", "smiles", "formula", "cid"],
                                  "description": "как трактовать значение, по умолчанию name"},
                },
                "required": ["name"],
            },
            handler=compound_lookup,
        ),
        Tool(
            name="make_document",
            description="Собирает готовый файл Word (.docx) или Excel (.xlsx) и возвращает "
            "ссылку на скачивание. Используйте, когда результат работы — документ: "
            "протокол испытаний, техническое задание, карта партии, раздел отчёта, "
            "выгрузка расчётов. Блоки: heading (заголовок), text (абзац), list (список), "
            "table (columns + rows, где каждая строка — объект с ключами из columns).",
            parameters={
                "type": "object",
                "properties": {
                    "format": {"type": "string", "enum": ["docx", "xlsx"], "description": "по умолчанию docx"},
                    "title": {"type": "string", "description": "заголовок и имя файла"},
                    "blocks": {
                        "type": "array",
                        "description": "содержимое: [{type:'heading'|'text'|'list'|'table', …}]",
                        "items": {"type": "object"},
                    },
                },
                "required": ["title", "blocks"],
            },
            handler=make_document,
        ),
        Tool(
            name="get_recipe",
            description="Методики института: проверенные порядки инженерного расчёта — "
            "что считать, в какой последовательности, каким кодом, какие пороги означают "
            "«так делать нельзя» и где типовые ошибки. Вызывайте БЕЗ аргументов, чтобы "
            "увидеть список, и с name — чтобы прочитать нужную ПЕРЕД тем, как писать код "
            "в run_calc: методика даёт правильные сигнатуры функций сразу, без угадывания, "
            "и предупреждает об ошибках, которые дают правдоподобный неверный результат.",
            parameters={
                "type": "object",
                "properties": {
                    "name": {"type": "string",
                             "description": "имя методики из списка; без него вернётся список"},
                },
            },
            handler=get_recipe,
        ),
        Tool(
            name="notify",
            description="Отправляет пользователю на его почту письмо с результатом. "
            "Используйте, когда пользователь просит прислать итог или отчёт письмом, "
            "либо когда длинная работа завершена и он не ждёт у экрана. Письмо уходит "
            "только самому пользователю, другим адресатам отправить нельзя. Не более "
            "10 писем в час; не отправляйте письмо без просьбы или содержательной причины.",
            parameters={
                "type": "object",
                "properties": {
                    "subject": {"type": "string", "description": "тема письма"},
                    "body": {"type": "string", "description": "текст, обычным языком, без разметки"},
                },
                "required": ["subject", "body"],
            },
            handler=notify,
        ),
    ]

    # run_calc регистрируется только при установленной песочнице: инструмент,
    # который всегда возвращает «не установлено», хуже отсутствующего —
    # модель будет строить на нём планы.
    if compute_sandbox.sandbox_available():
        tools.append(
            Tool(
                name="run_calc",
                description=(
                    "Выполняет инженерный расчёт: вы пишете короткий скрипт на Python, "
                    "платформа исполняет его и возвращает напечатанное, значение переменной "
                    "result и созданные файлы (графики .png сохраняйте через "
                    "matplotlib.savefig). Доступны numpy, scipy, pandas, sympy, pint "
                    "(единицы), matplotlib, scikit-learn, statsmodels, rdkit (химия), "
                    "thermo/chemicals/fluids/ht (теплофизика, гидравлика, теплообмен), "
                    "CoolProp. Сети в расчёте нет; данные участника подключайте через "
                    "data_files.\n\n"
                    "ОБЯЗАТЕЛЬНО используйте для любых чисел: площадей, расходов, "
                    "стехиометрии, выходов, подгонки моделей — НЕ считайте в уме и не "
                    "берите числа из памяти.\n\n"
                    "ЕСЛИ ВЕРНУЛСЯ ok:false — НЕ НАЗЫВАЙТЕ НИКАКИХ ЧИСЕЛ. Прочитайте "
                    "error, исправьте код и вызовите инструмент снова. Не вышло за "
                    "две-три попытки — прямо скажите человеку, что расчёт не выполнен и "
                    "что именно не получилось. Подставлять правдоподобные значения вместо "
                    "неудавшегося расчёта запрещено: уверенно поданное неверное число "
                    "хуже отсутствия ответа — по нему закажут не тот аппарат.\n\n"
                    "Сигнатуры, в которых чаще всего ошибаются (проверено на этом сервере, "
                    "не угадывайте их и не тратьте вызов на inspect):\n"
                    "  ht.LMTD(Thi, Tho, Tci, Tco, counterflow=True)\n"
                    "  ht.F_LMTD_Fakheri(Thi, Tho, Tci, Tco, shells=1)\n"
                    "  ht.effectiveness_NTU_method(mh, mc, Cph, Cpc, subtype='counterflow', "
                    "Thi=, Tho=, Tci=, Tco=, UA=)\n"
                    "  fluids.Reynolds(V, D, rho=, mu=)   fluids.friction_factor(Re, eD=0.0)\n"
                    "  fluids.K_from_f(fd, L, D)          fluids.dP_from_K(K, rho, V)\n"
                    "  fluids.nearest_pipe(Do=, Di=, NPS=, schedule='40')\n"
                    "  thermo.Chemical(ID, T=298.15, P=101325)  thermo.Mixture(IDs, zs=|ws=)\n"
                    "  chemicals.MW(ID), chemicals.Tb(CAS), chemicals.Tc(CAS), "
                    "chemicals.CAS_from_any(ID)\n"
                    "  CoolProp.CoolProp.PropsSI('D','T',T_K,'P',Pa,'Water')  — "
                    "температура ТОЛЬКО в кельвинах\n\n"
                    "Результат показывайте человеку компонентом calc_trace — инженер "
                    "должен видеть формулу и подставленные значения, а не только ответ."
                ),
                parameters={
                    "type": "object",
                    "properties": {
                        "code": {
                            "type": "string",
                            "description": "скрипт на Python; печатайте пояснения через print(), "
                            "итог кладите в переменную result (словарь с числами и единицами)",
                        },
                        "data_files": {
                            "type": "array",
                            "items": {"type": "string"},
                            "description": "id вложений (att-…), которые надо положить рядом "
                            "со скриптом; открываются обычным open()/pandas.read_excel()",
                        },
                    },
                    "required": ["code"],
                },
                handler=run_calc,
            )
        )
    return tools
